#!/usr/bin/env python3
"""Export the bookmark workbook into agent-friendly CSV and JSON files."""

from __future__ import annotations

import argparse
import csv
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCE_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "129a0kNYoIIGTySJiPQkzAw_-CVQ8VFlV3QU6FHlIXh0/edit"
)


def normalize(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def extract_table(sheet: Any, header_row: int = 4) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    headers: list[str] | None = None
    records: list[dict[str, Any]] = []

    for row_number, row in enumerate(rows, start=1):
        values = [normalize(value) for value in row]
        if row_number == header_row:
            headers = [str(value).strip() if value is not None else "" for value in values]
            continue
        if row_number <= header_row or not headers:
            continue
        if not any(value not in (None, "") for value in values):
            continue

        record = {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers)
            if header
        }
        records.append(record)

    return records


def write_csv(path: Path, records: list[dict[str, Any]]) -> None:
    if not records:
        raise ValueError(f"No records available for {path.name}")
    fieldnames = list(records[0].keys())
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(records)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)

    all_bookmarks = extract_table(workbook["All Bookmarks"])
    priority_queue = extract_table(workbook["Priority Queue"])
    manually_shared = extract_table(workbook["Manually Shared"])

    indexes = [record["#"] for record in all_bookmarks]
    tweet_urls = [record["Tweet URL"] for record in all_bookmarks]
    if len(all_bookmarks) != 250:
        raise ValueError(f"Expected 250 bookmarks, found {len(all_bookmarks)}")
    if indexes != list(range(1, 251)):
        raise ValueError("Bookmark indexes are not continuous from 1 through 250")
    if len(tweet_urls) != len(set(tweet_urls)):
        raise ValueError("Tweet URLs are not unique")

    write_csv(args.output_dir / "all-bookmarks.csv", all_bookmarks)
    write_csv(args.output_dir / "priority-queue.csv", priority_queue)
    write_csv(args.output_dir / "manually-shared.csv", manually_shared)

    usefulness_counts: dict[str, int] = {}
    for record in all_bookmarks:
        key = str(record.get("Usefulness") or "unclassified")
        usefulness_counts[key] = usefulness_counts.get(key, 0) + 1

    catalog = {
        "schema_version": 1,
        "captured_on": "2026-09-17",
        "source": {
            "title": "Twitter Bookmarks — Complete Catalog (250)",
            "url": SOURCE_URL,
        },
        "counts": {
            "all_bookmarks": len(all_bookmarks),
            "priority_queue": len(priority_queue),
            "manually_shared": len(manually_shared),
            "usefulness": usefulness_counts,
        },
        "all_bookmarks": all_bookmarks,
        "priority_queue": priority_queue,
        "manually_shared": manually_shared,
    }
    (args.output_dir / "catalog.json").write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(json.dumps(catalog["counts"], sort_keys=True))


if __name__ == "__main__":
    main()
